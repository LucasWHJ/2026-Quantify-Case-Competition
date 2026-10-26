"""Read case exhibits, calculate metrics, and export aggregate evidence.

No row-level incident data are exported. Source monetary values in Exhibit A
are CAD thousands; this module normalizes them to CAD.
"""
from pathlib import Path
import argparse
import json
import os
import hashlib
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from .pilot import Pilot

INCIDENT_COLUMNS = ["incident_id", "description", "incident_date", "industry",
                    "failure_type", "human_oversight", "estimated_insured_loss_cad"]
ROOT = Path(__file__).resolve().parents[1]


def locate_workbook(path=None):
    selected = path or os.environ.get("QUANTIFY_WORKBOOK") or ROOT / "data/raw/exhibits.xlsx"
    selected = Path(selected).expanduser()
    if not selected.is_file():
        raise FileNotFoundError(
            "The organizer workbook is not bundled. Place an authorized copy at "
            "data/raw/exhibits.xlsx or set QUANTIFY_WORKBOOK. See data/README.md. "
            "The pilot notebook and checked-in aggregate results work without it."
        )
    return selected


def load_exhibits(path=None):
    selected = locate_workbook(path)
    book = load_workbook(selected, data_only=True, read_only=True)
    if not {"Exhibit A", "Exhibit B"}.issubset(book.sheetnames):
        raise ValueError("Expected the combined workbook with Exhibit A and Exhibit B")
    sheet = book["Exhibit A"]
    records = []
    for anchor in (3, 13, 23, 33, 43):
        for col in range(2, 7):
            record = {
                "line": sheet.cell(anchor, 1).value,
                "year": int(sheet.cell(anchor + 1, col).value),
                "exposure": float(sheet.cell(anchor + 2, col).value),
                "earned_premium_cad": float(sheet.cell(anchor + 4, col).value) * 1000,
                "claim_count": int(sheet.cell(anchor + 5, col).value),
                "incurred_losses_cad": float(sheet.cell(anchor + 6, col).value) * 1000,
                "operating_expenses_cad": float(sheet.cell(anchor + 8, col).value) * 1000,
            }
            records.append(record)
    financial = pd.DataFrame(records)
    incident_records = []
    for row in book["Exhibit B"].iter_rows(values_only=True):
        if isinstance(row[0], str) and row[0].startswith("AI-"):
            incident_records.append(row[:7])
    incidents = pd.DataFrame(incident_records, columns=INCIDENT_COLUMNS)
    book.close()
    if incidents.empty or incidents[INCIDENT_COLUMNS].isna().any().any():
        raise ValueError("Empty incident sample or missing required fields")
    if incidents.incident_id.duplicated().any():
        raise ValueError("Duplicate incident IDs require explicit review")
    if not incidents.human_oversight.isin(["Yes", "No"]).all():
        raise ValueError("Unexpected oversight category")
    incidents["estimated_insured_loss_cad"] = pd.to_numeric(incidents.estimated_insured_loss_cad)
    if (incidents.estimated_insured_loss_cad <= 0).any():
        raise ValueError("Log-loss analysis requires strictly positive estimates")
    if (financial.exposure <= 0).any() or (financial.earned_premium_cad <= 0).any():
        raise ValueError("Positive exposure and premium required")
    return financial, incidents


def financial_metrics(financial):
    f = financial.copy()
    f["claim_frequency"] = f.claim_count / f.exposure
    f["average_severity_cad"] = f.incurred_losses_cad / f.claim_count
    f["pure_premium_cad"] = f.incurred_losses_cad / f.exposure
    f["premium_per_exposure_cad"] = f.earned_premium_cad / f.exposure
    f["loss_ratio"] = f.incurred_losses_cad / f.earned_premium_cad
    f["expense_ratio"] = f.operating_expenses_cad / f.earned_premium_cad
    f["combined_ratio"] = f.loss_ratio + f.expense_ratio
    f["underwriting_margin"] = 1 - f.combined_ratio
    f["underwriting_profit_cad"] = f.earned_premium_cad - f.incurred_losses_cad - f.operating_expenses_cad
    totals = f.groupby("year")[["earned_premium_cad", "incurred_losses_cad",
                              "operating_expenses_cad", "underwriting_profit_cad"]].sum()
    totals["combined_ratio"] = (totals.incurred_losses_cad + totals.operating_expenses_cad) / totals.earned_premium_cad
    # Do not sum heterogeneous exposure units or claim frequencies across lines.
    return f, totals.reset_index()


def severity_summary(incidents, groups):
    return incidents.groupby(groups, observed=True).estimated_insured_loss_cad.agg(
        n="count", mean_cad="mean", median_cad="median", minimum_cad="min",
        maximum_cad="max", total_cad="sum").reset_index()


def candidate_subset(incidents):
    """Descriptive candidate group, BEFORE legal/coverage eligibility mapping."""
    return incidents[(incidents.failure_type.isin(["Hallucinated Output", "Misinformation"]))
                     & (incidents.human_oversight == "Yes")].copy()


def design_matrix(data, columns):
    parts = [np.ones((len(data), 1))]
    names = ["Intercept"]
    for column in columns:
        categories = sorted(data[column].dropna().unique())
        for category in categories[1:]:
            parts.append((data[column].to_numpy() == category).astype(float)[:, None])
            names.append(f"{column}={category}")
    return np.column_stack(parts), names


def ols(data, columns):
    y = np.log(data.estimated_insured_loss_cad.to_numpy(dtype=float))
    x, names = design_matrix(data, columns)
    coefficients = np.linalg.lstsq(x, y, rcond=None)[0]
    residuals = y - x @ coefficients
    sse = float(residuals @ residuals)
    sst = float(np.sum((y - y.mean()) ** 2))
    rank = int(np.linalg.matrix_rank(x))
    return {"columns": columns, "terms": names, "coefficients": coefficients.tolist(),
            "n": len(y), "rank": rank, "sse": sse, "sst": sst,
            "r_squared": 1 - sse / sst,
            "adjusted_r_squared": 1 - (sse / (len(y) - rank)) / (sst / (len(y) - 1))}


def sequential_shares(data, order):
    baseline = ols(data, [])
    previous = baseline["sse"]
    records = []
    for index, term in enumerate(order):
        current = ols(data, order[:index + 1])
        records.append({"order": ", ".join(order), "term": term,
                        "share_of_total_log_variance": (previous - current["sse"]) / baseline["sst"]})
        previous = current["sse"]
    records.append({"order": ", ".join(order), "term": "Residual",
                    "share_of_total_log_variance": previous / baseline["sst"]})
    return pd.DataFrame(records)


def model_diagnostics(data):
    order = ["industry", "failure_type", "human_oversight"]
    model = ols(data, order)
    independent = []
    for term in order:
        reduced = ols(data, [c for c in order if c != term])
        independent.append({"term": term, "unique_share_of_total_log_variance":
                            (reduced["sse"] - model["sse"]) / model["sst"]})
    x, _ = design_matrix(data, order)
    y = np.log(data.estimated_insured_loss_cad.to_numpy(dtype=float))
    indices = np.random.default_rng(42).permutation(len(data))
    predictions = np.zeros(len(data))
    for test in np.array_split(indices, 5):
        train = np.setdiff1d(np.arange(len(data)), test)
        if np.linalg.matrix_rank(x[train]) < x.shape[1]:
            raise ValueError("A fold loses category support; revise validation design")
        predictions[test] = x[test] @ np.linalg.lstsq(x[train], y[train], rcond=None)[0]
    model["random_five_fold_log_r_squared"] = 1 - float(np.sum((y - predictions) ** 2)) / model["sst"]
    model["random_five_fold_mean_absolute_log_error"] = float(np.mean(np.abs(y - predictions)))
    model["oversight_log_coefficient"] = model["coefficients"][model["terms"].index("human_oversight=Yes")]
    model["oversight_conditional_geometric_mean_ratio"] = float(np.exp(model["oversight_log_coefficient"]))
    return model, pd.DataFrame(independent)


def build_evidence(path=None, output_dir=None):
    source = locate_workbook(path)
    financial, incidents = load_exhibits(source)
    metrics, yearly = financial_metrics(financial)
    model, unique = model_diagnostics(incidents)
    sequential = pd.concat([
        sequential_shares(incidents, ["industry", "failure_type", "human_oversight"]),
        sequential_shares(incidents, ["failure_type", "industry", "human_oversight"])
    ], ignore_index=True)
    candidate = candidate_subset(incidents)
    aggregate = {
        "source_workbook_sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
        "source_description": "Organizer-supplied combined Exhibits workbook, Quantify 2026",
        "financial_source_units": "CAD thousands, normalized to CAD",
        "incident_loss_measure": "Estimated insured loss in CAD; external incidents, not TSI claims",
        "n_incidents": len(incidents),
        "mean_estimated_loss_cad": float(incidents.estimated_insured_loss_cad.mean()),
        "median_estimated_loss_cad": float(incidents.estimated_insured_loss_cad.median()),
        "maximum_estimated_loss_cad": int(incidents.estimated_insured_loss_cad.max()),
        "top_nine_share": float(incidents.nlargest(9, "estimated_insured_loss_cad").estimated_insured_loss_cad.sum()
                                 / incidents.estimated_insured_loss_cad.sum()),
        "candidate_n": len(candidate),
        "candidate_mean_estimated_loss_cad": float(candidate.estimated_insured_loss_cad.mean()),
        "candidate_maximum_estimated_loss_cad": int(candidate.estimated_insured_loss_cad.max()),
        "model": model,
        "illustrative_pilot": Pilot().result(),
    }
    tables = {
        "company_annual": yearly,
        "line_metrics": metrics[["line", "year", "claim_frequency", "average_severity_cad", "pure_premium_cad",
                                 "premium_per_exposure_cad", "loss_ratio", "expense_ratio", "combined_ratio",
                                 "underwriting_margin", "underwriting_profit_cad"]],
        "severity_by_failure": severity_summary(incidents, ["failure_type"]),
        "severity_by_oversight": severity_summary(incidents, ["human_oversight"]),
        "severity_by_failure_and_oversight": severity_summary(incidents, ["failure_type", "human_oversight"]),
        "candidate_subset": severity_summary(candidate, ["failure_type"]),
        "anova_order_sensitivity": sequential,
        "unique_model_contributions": unique,
        "failure_industry_counts": pd.crosstab(incidents.failure_type, incidents.industry).reset_index(),
    }
    if output_dir is not None:
        output = Path(output_dir)
        output.mkdir(parents=True, exist_ok=True)
        for name, table in tables.items():
            table.to_csv(output / f"{name}.csv", index=False)
        (output / "key_results.json").write_text(json.dumps(aggregate, indent=2) + "\n", encoding="utf-8")
    return aggregate, tables


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", default=None)
    parser.add_argument("--output", type=Path, default=ROOT / "data/processed")
    parser.add_argument("--figures", action="store_true")
    args = parser.parse_args()
    _, tables = build_evidence(args.workbook, args.output)
    if args.figures:
        from .plots import generate_figures
        generate_figures(tables, ROOT / "figures")
    print("Aggregate evidence regenerated. No incident-level data exported.")


if __name__ == "__main__":
    main()
